# -*- coding: utf-8 -*-
"""
Generate a Java developer year-end review PPTX from the provided template.

It uses a simple, template-agnostic strategy:
- Open the template .pptx
- Traverse all slides/shapes that have text
- Replace placeholder tokens if they exist (e.g. {{TITLE}}, {{NAME}})
- If no tokens exist, it can also inject an "Appendix" slide with the full outline
  so the output is guaranteed to contain the intended content.

This is designed to work even when the template is not built with PowerPoint placeholders.
"""

from __future__ import annotations

import argparse
import datetime as _dt
import os
from pathlib import Path
from typing import Dict, Iterable

from pptx import Presentation
from pptx.util import Pt


DEFAULT_TEMPLATE = "附件1年终述职模板20260128.pptx"
DEFAULT_OUTPUT = "outputs/2025年终述职-Java程序员-示例.pptx"


def _iter_text_runs(shape) -> Iterable:
    """Yield (run) objects for a shape that contains text."""
    if not getattr(shape, "has_text_frame", False):
        return
    tf = shape.text_frame
    for p in tf.paragraphs:
        for run in p.runs:
            yield run


def replace_tokens(prs: Presentation, mapping: Dict[str, str]) -> int:
    """Replace tokens across all slides. Returns number of replacements."""
    count = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue

            # Replace at run level to preserve formatting as much as possible
            for run in _iter_text_runs(shape):
                if not run.text:
                    continue
                original = run.text
                new_text = original
                for k, v in mapping.items():
                    if k in new_text:
                        new_text = new_text.replace(k, v)
                if new_text != original:
                    run.text = new_text
                    count += 1
    return count


def add_appendix_slide(prs: Presentation, title: str, lines: Iterable[str]) -> None:
    """Always adds a final slide containing the full outline."""
    # Use a generic layout if available; fallback to first layout.
    layout = prs.slide_layouts[1] if len(prs.slide_layouts) > 1 else prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)

    # Try to set title/body if the layout supports it
    if slide.shapes.title:
        slide.shapes.title.text = title

    body = None
    for shp in slide.shapes:
        if shp.has_text_frame and shp is not slide.shapes.title:
            body = shp
            break

    if body is None:
        # Add a textbox if the chosen layout has no body
        left = top = Pt(24)
        width = Pt(720)
        height = Pt(540)
        body = slide.shapes.add_textbox(left, top, width, height)

    tf = body.text_frame
    tf.clear()

    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.text = line
        p.font.size = Pt(16)


def build_default_content() -> Dict[str, str]:
    """Token mapping for templates that include placeholders."""
    year = str(_dt.date.today().year - 1)  # assume "last year review"
    today = _dt.date.today().strftime("%Y-%m-%d")

    return {
        "{{TITLE}}": f"{year} 年终述职 / 工作总结（Java 程序员）",
        "{{SUBTITLE}}": "后端研发 · 交付与稳定性 · 技术沉淀",
        "{{NAME}}": "（请替换姓名）",
        "{{DEPT}}": "（请替换部门/团队）",
        "{{DATE}}": today,
        "{{AGENDA}}": "1. 工作概览\n2. 关键项目\n3. 技术建设\n4. 稳定性与质量\n5. 成长与协作\n6. 复盘与计划",
    }


def build_appendix_outline() -> list[str]:
    """A safe, non-sensitive, professional outline to embed even if template has no tokens."""
    return [
        "一、年度工作概览",
        "• 负责：核心业务接口与服务治理（用户中心/订单/支付对账等泛化场景）",
        "• 关键词：Spring Boot、MyBatis、Redis、MQ、MySQL、Docker/K8s、Prometheus/Grafana",
        "• 交付：按期完成版本迭代与需求交付，支撑业务增长",
        "",
        "二、关键项目与成果（示例）",
        "1) 核心链路性能优化",
        "   • 问题：高峰期接口 P95 延迟偏高、慢 SQL/锁等待",
        "   • 动作：SQL 索引与执行计划优化、缓存分层、异步化、批量化",
        "   • 结果：P95 从 ~800ms 降到 ~220ms；峰值 QPS 提升 ~2.5x（示例值，可改）",
        "",
        "2) 稳定性与可观测性建设",
        "   • 动作：统一日志规范、TraceId 贯通、关键指标面板、告警分级与值班SOP",
        "   • 结果：故障定位时间显著下降；核心服务可用性目标 99.9%+（按实际替换）",
        "",
        "3) 发布与交付效率提升",
        "   • 动作：CI（单测/静态扫描）+ CD（灰度/回滚）流程固化；版本变更清单",
        "   • 结果：发布失败率下降；回滚更可控；交付周期缩短（按实际替换）",
        "",
        "三、技术沉淀",
        "• 公共组件：统一异常与返回体、参数校验、幂等/限流、分页与导出封装",
        "• 代码质量：重构历史包袱模块；降低圈复杂度；提升单测覆盖（按实际）",
        "• 数据治理：热点Key治理、缓存穿透/击穿/雪崩预案",
        "",
        "四、协作与影响力",
        "• 需求评审：推动边界与验收标准明确；提前识别风险",
        "• Code Review：落实规范、分享最佳实践；带新人上手核心模块",
        "• 跨团队：对齐接口契约与灰度策略，减少联调成本",
        "",
        "五、不足与复盘",
        "• 对复杂问题的系统性拆解仍需加强（提前压测/容量评估）",
        "• 业务理解深度不均衡，需加强数据化分析与闭环",
        "• 文档沉淀有待提升（架构图、时序图、故障复盘模板）",
        "",
        "六、明年计划（示例）",
        "• 架构：推进服务分层与域模型梳理；关键链路解耦",
        "• 稳定性：容量模型、压测平台化、演练常态化",
        "• 效率：标准化脚手架、自动化回归、发布策略优化",
        "• 成长：深入 JVM/GC、并发、分布式一致性；输出分享与沉淀",
    ]


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", default=DEFAULT_TEMPLATE, help="Path to template pptx")
    parser.add_argument("--output", default=DEFAULT_OUTPUT, help="Path to output pptx")
    args = parser.parse_args()

    template_path = Path(args.template)
    output_path = Path(args.output)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    if not template_path.exists():
        raise FileNotFoundError(f"Template not found: {template_path}")

    prs = Presentation(str(template_path))

    mapping = build_default_content()
    replaced = replace_tokens(prs, mapping)

    # Always append a final slide so the deck definitely contains the intended content,
    # even when the template has no placeholders.
    appendix_title = "附录：Java 程序员年终述职示例内容（可直接替换为你的真实数据）"
    add_appendix_slide(prs, appendix_title, build_appendix_outline())

    prs.save(str(output_path))

    print(f"Generated: {output_path}")
    print(f"Token replacements applied: {replaced} (only if template contains tokens like {{TITLE}})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
